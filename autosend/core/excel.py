"""Read the school's Excel sheet into plain rows. The Excel file stays the master copy."""
from __future__ import annotations

import os
import shutil
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from .phone import normalize_phone


class ExcelError(Exception):
    """Human-readable problem with the spreadsheet."""


@dataclass
class RosterRow:
    key: str
    row_no: int
    values: dict[str, str]
    phone: str | None
    phone_error: str | None


@dataclass
class Roster:
    path: str
    columns: list[str]
    rows: list[RosterRow]
    key_col: str | None
    phone_col: str | None
    name_col: str | None
    mtime: float = 0.0
    class_col: str | None = None


def _cell_to_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return f"{v:.10g}"
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y") if (v.hour, v.minute, v.second) == (0, 0, 0) else v.strftime("%d/%m/%Y %H:%M")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


def _find(columns: list[str], hints: tuple[str, ...]) -> str | None:
    for hint in hints:
        for c in columns:
            if hint in c.lower():
                return c
    return None


def detect_columns(columns: list[str]) -> dict[str, str | None]:
    return {
        "phone": _find(columns, ("whatsapp", "phone", "mobile", "contact", "tel")),
        "key": _find(columns, ("admission", "adm", "reg", "id")),
        "name": _find(columns, ("name",)),
        "class": _find(columns, ("class", "grade")),
    }


def _open_readable(path: str) -> tuple[str, str | None]:
    """Return (path_to_read, temp_copy_to_delete). Copies when Excel holds the file."""
    try:
        with open(path, "rb"):
            pass
        return path, None
    except PermissionError:
        pass
    fd, tmp = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    try:
        shutil.copyfile(path, tmp)
    except OSError as e:
        os.unlink(tmp)
        raise ExcelError("The Excel file is busy right now. It will be read again in a moment.") from e
    return tmp, tmp


def read_roster(path: str, key_col: str | None = None, phone_col: str | None = None,
                default_cc: str = "94") -> Roster:
    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException

    p = Path(path)
    if not p.exists():
        raise ExcelError("The Excel file could not be found. Was it moved or renamed?")
    if p.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ExcelError("Please choose an Excel file that ends in .xlsx")

    read_path, tmp = _open_readable(str(p))
    wb = None
    try:
        try:
            wb = openpyxl.load_workbook(read_path, read_only=True, data_only=True)
        except (InvalidFileException, KeyError, OSError, ValueError) as e:
            raise ExcelError("This Excel file could not be opened. Is it saved as .xlsx?") from e
        except Exception as e:  # zipfile.BadZipFile etc. while Excel is mid-save
            raise ExcelError("The Excel file is being saved. It will be read again in a moment.") from e
        ws = wb.active or wb.worksheets[0]

        it = ws.iter_rows(values_only=True)
        header: list[str] | None = None
        raw_rows: list[tuple[int, tuple]] = []
        for i, row in enumerate(it, start=1):
            if header is None:
                if any(_cell_to_str(c) for c in row):
                    header = [_cell_to_str(c) for c in row]
                    header_row = i
                continue
            raw_rows.append((i, row))
        if header is None:
            raise ExcelError("The Excel file is empty. Add a header row and your students.")

        # unique, non-empty headers
        columns: list[str] = []
        col_index: list[int] = []
        seen: set[str] = set()
        for idx, h in enumerate(header):
            if not h:
                continue
            name = h
            n = 2
            while name.lower() in seen:
                name = f"{h} ({n})"
                n += 1
            seen.add(name.lower())
            columns.append(name)
            col_index.append(idx)

        det = detect_columns(columns)
        phone_c = phone_col if phone_col in columns else det["phone"]
        key_c = key_col if key_col in columns else det["key"]
        name_c = det["name"]
        class_c = det["class"]

        rows: list[RosterRow] = []
        used_keys: set[str] = set()
        for row_no, cells in raw_rows:
            values = {c: _cell_to_str(cells[i]) if i < len(cells) else "" for c, i in zip(columns, col_index)}
            if not any(values.values()):
                continue
            key = (values.get(key_c, "") if key_c else "").strip() or f"row{row_no}"
            base, n = key, 2
            while key in used_keys:
                key = f"{base}#{n}"
                n += 1
            used_keys.add(key)
            digits, err = normalize_phone(values.get(phone_c, "") if phone_c else "", default_cc)
            rows.append(RosterRow(key=key, row_no=row_no, values=values, phone=digits, phone_error=err))
        return Roster(path=str(p), columns=columns, rows=rows, key_col=key_c, phone_col=phone_c,
                      name_col=name_c, mtime=p.stat().st_mtime, class_col=class_c)
    finally:
        if wb is not None:
            wb.close()
        if tmp:
            try:
                os.unlink(tmp)
            except OSError:
                pass


def write_starter_workbook(path: str) -> None:
    """A tidy blank-ish sheet with the expected headers and three example rows (no phone numbers)."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.append(["Admission No.", "Name", "Class", "Whatsapp Number", "Outstanding Amount(Rs.)"])
    ws.append(["1001", "Example Student One", "Grade 5 A", "", 12000])
    ws.append(["1002", "Example Student Two", "Grade 5 B", "", 8500])
    ws.append(["1003", "Example Student Three", "Grade 6 A", "", 15000])
    fill = PatternFill("solid", fgColor="1FA855")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
        c.alignment = Alignment(vertical="center")
    for col, w in zip("ABCDE", (16, 34, 14, 22, 26)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    tips = wb.create_sheet("How to use")
    for line in (
        "Keep your students on the first sheet ('Students'). Keep the header row as it is.",
        "Type the parent's WhatsApp number in 'Whatsapp Number', e.g. +94 77 123 4567 or 0771234567.",
        "You can edit this file at any time. AutoSend notices when you save and updates itself.",
        "Delete the three example rows before you start.",
    ):
        tips.append([line])
    tips.column_dimensions["A"].width = 100
    wb.save(path)
