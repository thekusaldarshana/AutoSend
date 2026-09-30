"""Write student changes back into the client's Excel file, safely.

- A backup copy is kept (in the app's data folder) before the first change of each session.
- Rows are checked before they're touched, so a file that changed underneath us is never overwritten blindly.
- The file is saved to a temporary name and swapped in, so a crash can't leave a half-written workbook.
- If Excel has the file open, Windows refuses the swap; we report that plainly.
"""
from __future__ import annotations

import os
import re
import shutil
import time
from copy import copy
from pathlib import Path

from . import paths
from .excel import ExcelError, _cell_to_str

_backed_up: set[str] = set()


class ExcelBusy(ExcelError):
    """The file is open in Excel (or otherwise locked)."""


class ExcelChanged(ExcelError):
    """The row we meant to edit is no longer where we expected."""


def backup(path: str) -> Path | None:
    if path in _backed_up or not os.path.exists(path):
        return None
    d = paths.sub_dir("backups")
    dst = d / f"{Path(path).stem} {time.strftime('%Y-%m-%d %H-%M-%S')}{Path(path).suffix}"
    try:
        shutil.copyfile(path, dst)
    except OSError:
        return None
    _backed_up.add(path)
    # keep the 20 newest
    files = sorted(d.glob("*"), key=lambda p: p.stat().st_mtime, reverse=True)
    for old in files[20:]:
        try:
            old.unlink()
        except OSError:
            pass
    return dst


def _load(path: str):
    import openpyxl
    if not os.path.exists(path):
        raise ExcelError("The Excel file could not be found.")
    try:
        wb = openpyxl.load_workbook(path, keep_vba=path.lower().endswith(".xlsm"))
    except PermissionError as e:
        raise ExcelBusy("The Excel file is open. Close it in Excel, then try again.") from e
    except Exception as e:
        raise ExcelError("The Excel file could not be opened for editing.") from e
    return wb, wb.active


def _header(ws) -> tuple[int, dict[str, int]]:
    """(header_row, {column name: column index}) using the same naming rules as the reader."""
    for r in range(1, min(ws.max_row, 50) + 1):
        vals = [_cell_to_str(c.value) for c in ws[r]]
        if any(vals):
            cols: dict[str, int] = {}
            for i, h in enumerate(vals, start=1):
                if not h:
                    continue
                name, n = h, 2
                while name.lower() in {k.lower() for k in cols}:
                    name = f"{h} ({n})"
                    n += 1
                cols[name] = i
            return r, cols
    raise ExcelError("The Excel file has no header row.")


def _last_data_row(ws, header_row: int, ncols: int) -> int:
    r = ws.max_row
    while r > header_row:
        if any(_cell_to_str(ws.cell(r, c).value) for c in range(1, ncols + 1)):
            return r
        r -= 1
    return header_row


def _typed(value: str, keep_text: bool):
    v = value.strip()
    if keep_text or not v:
        return v or None
    if re.fullmatch(r"-?(0|[1-9]\d{0,14})", v):
        return int(v)
    if re.fullmatch(r"-?(0|[1-9]\d{0,14})\.\d{1,6}", v):
        return float(v)
    return v


def _save(wb, path: str) -> None:
    tmp = str(Path(path).with_name("~autosend-" + Path(path).name))
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    except PermissionError as e:
        raise ExcelBusy("The Excel file is open. Close it in Excel, then try again.") from e
    finally:
        wb.close()
        if os.path.exists(tmp):
            try:
                os.unlink(tmp)
            except OSError:
                pass


def _check_row(ws, row_no: int, key_idx: int | None, expect_key: str) -> None:
    if row_no < 1 or row_no > ws.max_row:
        raise ExcelChanged("The Excel file was changed. Please try again.")
    if key_idx is None:
        return
    actual = _cell_to_str(ws.cell(row_no, key_idx).value)
    want = expect_key.split("#")[0]
    if want.startswith("row") and not actual:
        return
    if actual != want:
        raise ExcelChanged("The Excel file was changed while you were editing. Please try again.")


def add_row(path: str, values: dict[str, str], text_cols: set[str]) -> None:
    backup(path)
    wb, ws = _load(path)
    hr, cols = _header(ws)
    last = _last_data_row(ws, hr, max(cols.values()))
    row = last + 1
    for name, idx in cols.items():
        if name in values:
            cell = ws.cell(row, idx, _typed(values[name], name in text_cols))
        else:
            cell = ws.cell(row, idx)
        src = ws.cell(last, idx)
        if last > hr and src.has_style:
            cell._style = copy(src._style)
    _save(wb, path)


def update_row(path: str, row_no: int, expect_key: str, key_col: str | None, values: dict[str, str],
               text_cols: set[str]) -> None:
    backup(path)
    wb, ws = _load(path)
    hr, cols = _header(ws)
    _check_row(ws, row_no, cols.get(key_col) if key_col else None, expect_key)
    for name, idx in cols.items():
        if name in values:
            ws.cell(row_no, idx).value = _typed(values[name], name in text_cols)
    _save(wb, path)


def delete_rows(path: str, rows: list[tuple[int, str]], key_col: str | None) -> None:
    backup(path)
    wb, ws = _load(path)
    hr, cols = _header(ws)
    kidx = cols.get(key_col) if key_col else None
    for row_no, key in rows:
        _check_row(ws, row_no, kidx, key)
    for row_no, _ in sorted(rows, reverse=True):
        ws.delete_rows(row_no)
    _save(wb, path)
